"""``voucher_template`` 新布局测试（纯逻辑；合成 25 列模板，**不提交真实模板**）。

覆盖 Issue #71 验收标准：

* ``app/export/`` 新增 ``voucher_template`` 布局适配器（25 列，**按列号定位**）；
* 模板装载：SHA-256 绑定 + 工作表存在 + 表头/列数与 sidecar 一致（fail closed）；
* 配置扩展：``[voucher]`` 凭证号模式、``[auxiliary]`` 命名辅助核算列（非法值 fail closed）；
* 两布局各自端到端 + 既有校验分支复用；
* sidecar 与代码表头不漂移（防回归）。

本文件不触库、不触网；本地运行：

    uv run pytest app/tests/export/test_voucher_layout.py -q -o addopts="" --noconftest
"""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.export.config import (
    VOUCHER_TEMPLATE_NAMED_COLUMNS,
    load_export_config,
)
from app.export.errors import (
    CODE_CONFIG_INVALID,
    CODE_CONFIG_MISSING_KEYS,
    CODE_TEMPLATE_HEADER_MISMATCH,
    CODE_UNSUPPORTED_LAYOUT,
    ExportConfigError,
    ExportTemplateError,
)
from app.export.mapping import ClaimInput, LineInput, build_entries
from app.export.template import (
    LAYOUT_VOUCHER_TEMPLATE,
    VOUCHER_TEMPLATE_HEADERS,
    WORKSHEET_ENTRY_ROWS,
    WORKSHEET_VOUCHER_TEMPLATE,
    load_template,
)
from app.export.validation import validate_export
from app.export.voucher_layout import (
    NAMED_AUXILIARY_HEADERS,
    voucher_template_rows,
)
from app.export.writer import render_entry_rows, render_export, render_voucher_template

from _helpers import build_config_toml, build_template_xlsx, make_config, write_config

REPO_ROOT = Path(__file__).resolve().parents[3]
SIDECAR = REPO_ROOT / "docs" / "export-templates" / "voucher-template-structure.json"
PERIOD = "2026-10"


def _line(**overrides: object) -> LineInput:
    base: dict[str, object] = {
        "line_id": 1,
        "expense_date": date(2026, 10, 5),
        "category_code": "CAT-TRV",
        "category_name": "差旅",
        "description": "出差打车",
        "amount": Decimal("100.00"),
        "tax_amount": Decimal("0.00"),
        "department_name": "财务部",
    }
    base.update(overrides)
    return LineInput(**base)  # type: ignore[arg-type]


def _claim(**overrides: object) -> ClaimInput:
    base: dict[str, object] = {
        "claim_id": 1,
        "claim_no": "EC-00000001",
        "currency": "CNY",
        "status": "paid",
        "total_amount": Decimal("100.00"),
        "employee_no": "E0001",
        "lines": (_line(),),
        "invoices": (),
        "payment_amount": Decimal("100.00"),
        "payment_date": date(2026, 10, 20),
    }
    base.update(overrides)
    return ClaimInput(**base)  # type: ignore[arg-type]


def _voucher_config(tmp_path: Path, **overrides: object):
    return make_config(tmp_path, layout=LAYOUT_VOUCHER_TEMPLATE, **overrides)


def _sheet(data: bytes, name: str = WORKSHEET_VOUCHER_TEMPLATE):
    return load_workbook(io.BytesIO(data))[name]


# --- 结构登记与 sidecar 不漂移 -------------------------------------------------


def test_voucher_headers_match_sidecar_and_named_columns() -> None:
    """代码登记的表头必须与 sidecar 逐列一致；命名列集合三处一致。"""
    assert len(VOUCHER_TEMPLATE_HEADERS) == 25
    sidecar = json.loads(SIDECAR.read_text(encoding="utf-8"))
    sheet = next(
        item for item in sidecar["worksheets"] if item["name"] == WORKSHEET_VOUCHER_TEMPLATE
    )
    assert [column["header"] for column in sheet["columns"]] == list(
        VOUCHER_TEMPLATE_HEADERS
    )
    assert [column["column"] for column in sheet["columns"]] == [
        "A", "B", "C", "D", "E", "F", "G", "H", "I", "J",
        "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T",
        "U", "V", "W", "X", "Y",
    ]
    assert NAMED_AUXILIARY_HEADERS == VOUCHER_TEMPLATE_HEADERS[10:20]
    assert VOUCHER_TEMPLATE_NAMED_COLUMNS == set(NAMED_AUXILIARY_HEADERS)


def test_sidecar_contains_structure_only() -> None:
    """sidecar 只允许结构：两张备查表只记表头与存在性，且无任何数据取值。"""
    sidecar = json.loads(SIDECAR.read_text(encoding="utf-8"))
    names = [item["name"] for item in sidecar["worksheets"]]
    assert names == [WORKSHEET_VOUCHER_TEMPLATE, "科目数据", "辅助核算项目数据"]
    # 备查表只登记列位置与（模板自带的）表头占位文本；不登记任何行取值。
    for item in sidecar["worksheets"]:
        for key in ("columns", "first_data_row"):
            assert key in item
        assert "rows" not in item
        for column in item["columns"]:
            assert "备查数据" in column["header"] or item["name"] == (
                WORKSHEET_VOUCHER_TEMPLATE
            )


# --- 模板装载（25 列） ---------------------------------------------------------


def test_template_loads_25_column_synthetic(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path)
    loaded = load_template(config)
    assert loaded.layout == LAYOUT_VOUCHER_TEMPLATE
    assert loaded.worksheet == WORKSHEET_VOUCHER_TEMPLATE
    assert loaded.column_count == 25
    assert loaded.sha256 == config.template.sha256


def test_template_column_order_change_fails_closed(tmp_path: Path) -> None:
    """列序变化 = 模板升版：即使列数相同也必须 fail closed（按列号定位的前提）。"""
    headers = list(VOUCHER_TEMPLATE_HEADERS)
    headers[0], headers[1] = headers[1], headers[0]
    template_path = tmp_path / "swapped.xlsx"
    sha = build_template_xlsx(
        template_path, headers=headers, sheet_name=WORKSHEET_VOUCHER_TEMPLATE
    )
    config = _voucher_config(
        tmp_path, headers=headers, template_path=str(template_path), template_sha=sha
    )
    with pytest.raises(ExportTemplateError) as excinfo:
        load_template(config)
    assert excinfo.value.code == CODE_TEMPLATE_HEADER_MISMATCH


def test_template_short_header_row_fails_closed(tmp_path: Path) -> None:
    headers = list(VOUCHER_TEMPLATE_HEADERS)[:24]
    template_path = tmp_path / "short.xlsx"
    sha = build_template_xlsx(
        template_path, headers=headers, sheet_name=WORKSHEET_VOUCHER_TEMPLATE
    )
    config = _voucher_config(
        tmp_path, headers=headers, template_path=str(template_path), template_sha=sha
    )
    with pytest.raises(ExportTemplateError) as excinfo:
        load_template(config)
    assert excinfo.value.code == CODE_TEMPLATE_HEADER_MISMATCH


def test_unknown_layout_still_fails_closed(tmp_path: Path) -> None:
    """未登记布局（如旧文档提到的 ``pz``）仍以 unsupported_layout 失败。"""
    import dataclasses

    config = _voucher_config(tmp_path)
    patched = dataclasses.replace(
        config, template=dataclasses.replace(config.template, layout="pz")
    )
    with pytest.raises(ExportTemplateError) as excinfo:
        load_template(patched)
    assert excinfo.value.code == CODE_UNSUPPORTED_LAYOUT


# --- 端到端渲染（按列号） ------------------------------------------------------


def test_voucher_template_end_to_end_columns(tmp_path: Path) -> None:
    config = _voucher_config(
        tmp_path, aux_named_columns={"部门": {"source": "department"}}
    )
    claims = [_claim()]
    entries = build_entries(claims, config, period=PERIOD)
    sheet = _sheet(render_voucher_template(entries, config))

    assert [cell.value for cell in sheet[1]] == list(VOUCHER_TEMPLATE_HEADERS)
    assert sheet.max_row == 3  # 表头 + 借方 + 贷方

    debit = sheet[2]
    assert debit[0].value.date() == date(2026, 10, 5)  # A 日期
    assert debit[0].number_format == "yyyy-mm-dd"
    assert debit[1].value == "记"  # B 凭证字
    assert debit[2].value == 1  # C 凭证号
    assert debit[3].value == 0  # D 附件数（无挂账发票）
    assert debit[4].value == 1  # E 分录序号
    assert debit[5].value == "EC-00000001 E0001 出差打车"  # F 摘要
    assert debit[6].value == "560101"  # G 科目代码
    assert debit[7].value == "管理费用-差旅费"  # H 科目名称
    assert debit[8].value == 100.0  # I 借方金额
    assert debit[9].value is None  # J 贷方金额
    assert debit[14].value == "财务部"  # O 部门（命名列，source=department）
    assert [debit[index].value for index in range(10, 14)] == [None] * 4  # K–N
    assert [debit[index].value for index in range(15, 20)] == [None] * 5  # P–T
    assert debit[20].value is None  # U 数量（安全默认留空）
    assert debit[21].value is None  # V 单价（安全默认留空）
    assert debit[22].value == 100.0  # W 原币金额
    assert debit[23].value == "RMB"  # X 币别
    assert debit[24].value == 1.0  # Y 汇率

    credit = sheet[3]
    assert credit[4].value == 2  # E 分录序号（同凭证内递增）
    assert credit[6].value == "100201"  # G 贷方科目
    assert credit[8].value is None  # I
    assert credit[9].value == 100.0  # J
    assert credit[14].value is None  # O 贷方分录不挂部门


def test_entry_sequence_increments_per_voucher(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path)
    lines = (
        _line(line_id=1, amount=Decimal("20.00")),
        _line(line_id=2, amount=Decimal("30.00")),
    )
    entries = build_entries(
        [_claim(total_amount=Decimal("50.00"), payment_amount=Decimal("50.00"), lines=lines)],
        config,
        period=PERIOD,
    )
    sheet = _sheet(render_voucher_template(entries, config))
    assert [sheet[row][4].value for row in range(2, 5)] == [1, 2, 3]


def test_number_start_offsets_voucher_numbers(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path, number_start=100)
    earlier = _claim(claim_id=3, claim_no="EC-00000003", lines=(_line(expense_date=date(2026, 10, 2)),))
    later = _claim(claim_id=5, claim_no="EC-00000005", lines=(_line(expense_date=date(2026, 10, 9)),))
    entries = build_entries([later, earlier], config, period=PERIOD)
    numbers = {entry.business_no: entry.voucher_no for entry in entries}
    assert numbers["EC-00000003"] == 100
    assert numbers["EC-00000005"] == 101
    sheet = _sheet(render_voucher_template(entries, config))
    assert sheet[2][2].value == 100


def test_number_mode_blank_writes_blank_voucher_no(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path, number_mode="blank")
    claims = [_claim()]
    entries = build_entries(claims, config, period=PERIOD)
    # 内部凭证号仍按顺序编号（校验/对账依赖它），仅写出时留空。
    assert {entry.voucher_no for entry in entries} == {1}
    sheet = _sheet(render_voucher_template(entries, config))
    assert sheet[2][2].value is None
    assert sheet[3][2].value is None
    assert validate_export(claims, entries, config, period=PERIOD) == []


def test_voucher_render_is_byte_reproducible(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path)
    entries = build_entries([_claim()], config, period=PERIOD)
    first = render_voucher_template(entries, config)
    second = render_voucher_template(build_entries([_claim()], config, period=PERIOD), config)
    assert first == second


def test_voucher_render_has_no_export_timestamp(tmp_path: Path) -> None:
    import zipfile

    config = _voucher_config(tmp_path)
    data = render_voucher_template(build_entries([_claim()], config, period=PERIOD), config)
    core = zipfile.ZipFile(io.BytesIO(data)).read("docProps/core.xml").decode()
    assert "1980-01-01" in core
    assert str(date.today().year) not in core


# --- 命名辅助核算列 ------------------------------------------------------------


def test_named_department_column(tmp_path: Path) -> None:
    config = _voucher_config(
        tmp_path,
        aux_named_columns={"部门": {"source": "department"}},
    )
    entries = build_entries([_claim()], config, period=PERIOD)
    assert entries[0].department_name == "财务部"
    rows = voucher_template_rows(entries, config)
    assert rows[0][14] == "财务部"  # O
    assert rows[1][14] is None  # 贷方分录无部门


def test_named_literal_column_and_injection_escape(tmp_path: Path) -> None:
    config = _voucher_config(
        tmp_path,
        aux_named_columns={
            "部门": {"source": "department"},
            "项目": {"source": "literal", "value": "=默认项目"},
        },
    )
    sheet = _sheet(render_voucher_template(build_entries([_claim()], config, period=PERIOD), config))
    # 配置文本同样过注入防护（§6.4）：危险前缀被单引号转义。
    assert sheet[2][13].value == "'=默认项目"


def test_named_columns_unregistered_stay_blank(tmp_path: Path) -> None:
    config = _voucher_config(tmp_path)
    assert config.auxiliary.named_columns == {}
    rows = voucher_template_rows(build_entries([_claim()], config, period=PERIOD), config)
    assert all(value is None for value in rows[0][10:20])


def test_named_columns_do_not_affect_old_layout(tmp_path: Path) -> None:
    """命名列只对新布局生效；旧布局的 31 列展开逐列不变。"""
    config = make_config(
        tmp_path,
        layout="entry_rows",
        aux_named_columns={"部门": {"source": "department"}},
    )
    entries = build_entries([_claim()], config, period=PERIOD)
    assert entries[0].cells()[18] is None  # 旧布局 S 列仍由 [auxiliary].mode 决定
    assert len(entries[0].cells()) == 31


# --- 布局切换（两布局各自端到端） ----------------------------------------------


def test_both_layouts_render_end_to_end(tmp_path: Path) -> None:
    vdir = tmp_path / "v"
    vdir.mkdir()
    ldir = tmp_path / "l"
    ldir.mkdir()

    voucher = _voucher_config(vdir, number_mode="blank")
    entries = build_entries([_claim()], voucher, period=PERIOD)
    voucher_sheet = _sheet(render_export(entries, voucher))
    assert [cell.value for cell in voucher_sheet[1]] == list(VOUCHER_TEMPLATE_HEADERS)

    legacy = make_config(ldir, layout="entry_rows")
    legacy_entries = build_entries([_claim()], legacy, period=PERIOD)
    legacy_sheet = _sheet(render_export(legacy_entries, legacy), WORKSHEET_ENTRY_ROWS)
    assert legacy_sheet.max_column == 31
    assert legacy_sheet[2][1].value == 1  # B 凭证号（旧布局仍按顺序号写出）


def test_render_export_dispatches_by_layout(tmp_path: Path) -> None:
    vdir = tmp_path / "v"
    vdir.mkdir()
    ldir = tmp_path / "l"
    ldir.mkdir()
    voucher = _voucher_config(vdir)
    legacy = make_config(ldir, layout="entry_rows")
    entries = build_entries([_claim()], voucher, period=PERIOD)
    assert render_export(entries, voucher) == render_voucher_template(entries, voucher)
    assert render_export(entries, legacy) == render_entry_rows(entries)


def test_render_export_rejects_unregistered_layout(tmp_path: Path) -> None:
    """``render_export`` 对未登记布局同样 fail closed（不静默降级）。"""
    import dataclasses

    config = _voucher_config(tmp_path)
    patched = dataclasses.replace(
        config, template=dataclasses.replace(config.template, layout="pz")
    )
    entries = build_entries([_claim()], config, period=PERIOD)
    with pytest.raises(ExportTemplateError) as excinfo:
        render_export(entries, patched)
    assert excinfo.value.code == CODE_UNSUPPORTED_LAYOUT


# --- 配置 fail closed ---------------------------------------------------------


def _write(tmp_path: Path, **overrides: object) -> Path:
    template_path = tmp_path / "t.xlsx"
    sha = build_template_xlsx(
        template_path, headers=VOUCHER_TEMPLATE_HEADERS, sheet_name=WORKSHEET_VOUCHER_TEMPLATE
    )
    overrides.setdefault("template_path", str(template_path))
    overrides.setdefault("template_sha", sha)
    overrides.setdefault("layout", LAYOUT_VOUCHER_TEMPLATE)
    overrides.setdefault("worksheet", WORKSHEET_VOUCHER_TEMPLATE)
    return write_config(tmp_path / "c.toml", build_config_toml(**overrides))  # type: ignore[arg-type]


def test_invalid_number_mode_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(_write(tmp_path, number_mode="auto"))
    assert excinfo.value.code == CODE_CONFIG_INVALID
    assert "voucher.number_mode" in excinfo.value.message


def test_number_start_zero_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(_write(tmp_path, number_start=0))
    assert excinfo.value.code == CODE_CONFIG_INVALID
    assert "voucher.number_start" in excinfo.value.message


def test_number_start_non_integer_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(_write(tmp_path, number_start="1"))
    assert excinfo.value.code == CODE_CONFIG_INVALID


def test_unknown_named_column_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(
            _write(tmp_path, aux_named_columns={"不存在的列": {"source": "blank"}})
        )
    assert excinfo.value.code == CODE_CONFIG_INVALID


def test_unknown_named_source_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(
            _write(tmp_path, aux_named_columns={"部门": {"source": "guess"}})
        )
    assert excinfo.value.code == CODE_CONFIG_INVALID


def test_named_literal_without_value_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(
            _write(tmp_path, aux_named_columns={"部门": {"source": "literal"}})
        )
    assert excinfo.value.code == CODE_CONFIG_MISSING_KEYS


def test_named_value_without_literal_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(
            _write(tmp_path, aux_named_columns={"部门": {"source": "blank", "value": "x"}})
        )
    assert excinfo.value.code == CODE_CONFIG_INVALID


def test_named_columns_non_table_fails_closed(tmp_path: Path) -> None:
    with pytest.raises(ExportConfigError) as excinfo:
        load_export_config(
            _write(tmp_path, extra_lines='[auxiliary.named_columns]\n部门 = "department"')
        )
    assert excinfo.value.code == CODE_CONFIG_INVALID
