"""写出层测试（纯逻辑；§6.4 注入防护 + §6.5 字节可复现）。

覆盖 Issue #53 验收标准第 6 条：

* 文本以 ``=``/``+``/``-``/``@`` 开头必须转义；纵深防御闸门可拒绝未转义文本；
* 只产 ``.xlsx``（ZIP 容器）、**不使用宏**（无 ``vbaProject.bin``）；
* 日期写日期型单元格、金额写数值单元格、整数列写整数；
* 同一批数据两次渲染字节一致（导出文件不含时间戳）。
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.export.errors import (
    CODE_CELL_INJECTION_UNESCAPED,
    ExportValidationError,
)
from app.export.mapping import EntryRow, build_entries
from app.export.template import ENTRY_ROWS_HEADERS, WORKSHEET_ENTRY_ROWS
from app.export.writer import (
    DANGEROUS_PREFIXES,
    XLSX_MEDIA_TYPE,
    escape_text,
    guard_no_unescaped,
    render_entry_rows,
)

from _helpers import make_config

PERIOD = "2026-10"


def _entry(**overrides: object) -> EntryRow:
    base: dict[str, object] = {
        "voucher_word": "记",
        "voucher_no": 1,
        "entry_date": date(2026, 10, 5),
        "summary": "EC-00000001 E0001 出差打车",
        "account_code": "560101",
        "account_name": "管理费用-差旅费",
        "currency_code": "RMB",
        "currency_name": "人民币",
        "original_amount": Decimal("100.00"),
        "exchange_rate": Decimal("1"),
        "debit": Decimal("100.00"),
        "credit": None,
        "preparer": "E0001",
        "business_no": "EC-00000001",
        "attachment_count": 2,
    }
    base.update(overrides)
    return EntryRow(**base)  # type: ignore[arg-type]


# --- 注入转义 -----------------------------------------------------------------


@pytest.mark.parametrize("prefix", DANGEROUS_PREFIXES)
def test_escape_text_prefixes_dangerous_values(prefix: str) -> None:
    assert escape_text(prefix + "SUM(A1)") == "'" + prefix + "SUM(A1)"


@pytest.mark.parametrize("value", ["差旅费", "EC-00000001", "", "合计"])
def test_escape_text_leaves_safe_values_untouched(value: str) -> None:
    assert escape_text(value) == value


def test_guard_rejects_unescaped_prefix() -> None:
    with pytest.raises(ExportValidationError) as excinfo:
        guard_no_unescaped([("ok", "=cmd|' /C calc'!A0", None)])
    assert excinfo.value.code == CODE_CELL_INJECTION_UNESCAPED
    assert excinfo.value.failures[0].rule == "fail_closed_12"


def test_guard_accepts_escaped_prefix() -> None:
    guard_no_unescaped([("ok", "'=1+1", 1)])


# --- 渲染往返 -----------------------------------------------------------------


def test_render_roundtrips_with_correct_types() -> None:
    data = render_entry_rows([_entry(), _entry(debit=None, credit=Decimal("100.00"))])

    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook[WORKSHEET_ENTRY_ROWS]

    assert [cell.value for cell in sheet[1]] == list(ENTRY_ROWS_HEADERS)
    assert sheet.max_row == 3

    debit = sheet[2]
    assert debit[0].value == "记"
    assert isinstance(debit[1].value, int) and debit[1].value == 1
    assert isinstance(debit[2].value, (date, datetime))
    assert debit[2].value.date() == date(2026, 10, 5)
    assert debit[2].number_format == "yyyy-mm-dd"
    assert debit[4].value == "560101"
    assert debit[6].value is None  # G 余额方向（工具列）
    assert debit[11].value is None  # L 数量
    assert debit[13].value == 100.0  # N 借方金额（数值）
    assert debit[14].value is None  # O 贷方金额
    assert debit[29].value == 2  # AD 附件数（整数）
    assert debit[30].value is None  # AE 检测错误情况

    credit = sheet[3]
    assert credit[13].value is None
    assert credit[14].value == 100.0


def test_render_escapes_injection_in_text_cells() -> None:
    data = render_entry_rows(
        [
            _entry(
                summary="=HYPERLINK(\"http://x\")",
                account_name="+danger",
            )
        ]
    )
    sheet = load_workbook(io.BytesIO(data))[WORKSHEET_ENTRY_ROWS]
    row = sheet[2]
    assert row[3].value == "'=HYPERLINK(\"http://x\")"
    assert row[5].value == "'+danger"


def test_render_is_byte_reproducible() -> None:
    entries = [_entry()]
    first = render_entry_rows(entries)
    second = render_entry_rows(list(entries))
    assert first == second


def test_render_is_deterministic_across_configs(tmp_path: Path) -> None:
    config = make_config(tmp_path)
    claims = []
    from app.export.mapping import ClaimInput, LineInput

    claims.append(
        ClaimInput(
            claim_id=1,
            claim_no="EC-00000001",
            currency="CNY",
            status="paid",
            total_amount=Decimal("100.00"),
            employee_no="E0001",
            lines=(
                LineInput(
                    line_id=1,
                    expense_date=date(2026, 10, 5),
                    category_code="CAT-TRV",
                    category_name="差旅",
                    description="出差打车",
                    amount=Decimal("100.00"),
                    tax_amount=Decimal("0.00"),
                ),
            ),
            payment_amount=Decimal("100.00"),
            payment_date=date(2026, 10, 20),
        )
    )
    first = render_entry_rows(build_entries(claims, config, period=PERIOD))
    second = render_entry_rows(build_entries(claims, config, period=PERIOD))
    assert first == second


def test_render_output_is_xlsx_zip_without_macros() -> None:
    data = render_entry_rows([_entry()])
    assert data.startswith(b"PK\x03\x04")
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert any(name.endswith("sheet1.xml") for name in names)
    assert not any("vbaProject" in name for name in names)
    assert XLSX_MEDIA_TYPE.endswith("spreadsheetml.sheet")


def test_render_has_no_export_timestamp_in_core_properties() -> None:
    data = render_entry_rows([_entry()])
    core = zipfile.ZipFile(io.BytesIO(data)).read("docProps/core.xml").decode()
    assert "1980-01-01" in core
    assert str(date.today().year) not in core


def test_render_empty_entries_writes_header_only() -> None:
    sheet = load_workbook(io.BytesIO(render_entry_rows([])))[WORKSHEET_ENTRY_ROWS]
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == list(ENTRY_ROWS_HEADERS)
