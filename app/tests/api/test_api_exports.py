"""Export endpoint tests (W1-11 acceptance point 8).

``GET /claims/{id}/export/pdf`` must return ``application/pdf`` bytes starting
with ``%PDF-``; ``GET /claims/{id}/export/xlsx`` must return the OOXML
spreadsheet content type and reload cleanly through ``openpyxl``.  Both must run
the same read authorization as ``GET /claims/{id}``.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.tests.api.conftest import World

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _draft_with_line(
    client: TestClient, headers: dict[str, str], world: World
) -> int:
    created = client.post("/claims", json={"currency": "CNY"}, headers=headers)
    assert created.status_code == 201, created.text
    claim_id = int(created.json()["id"])
    added = client.post(
        f"/claims/{claim_id}/lines",
        json={
            "expense_date": "2026-10-01",
            "category_id": world.category.id,
            "description": "高铁票",
            "amount": "100.10",
            "tax_amount": "3.00",
            "cost_center_id": world.cost_center.id,
        },
        headers=headers,
    )
    assert added.status_code == 201, added.text
    return claim_id


def test_export_pdf(
    client: TestClient, world: World, auth: Callable[..., dict[str, str]]
) -> None:
    headers = auth(world.claimant)
    claim_id = _draft_with_line(client, headers, world)
    response = client.get(f"/claims/{claim_id}/export/pdf", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF-")


def test_export_xlsx_reloadable(
    client: TestClient, world: World, auth: Callable[..., dict[str, str]]
) -> None:
    headers = auth(world.claimant)
    claim_id = _draft_with_line(client, headers, world)
    response = client.get(f"/claims/{claim_id}/export/xlsx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MIME)

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    # Header row + one detail row + total row.
    assert rows[0] == ("日期", "描述", "金额", "税额")
    assert rows[1][1] == "高铁票"
    assert rows[-1][0] == "合计"
    assert rows[-1][2] == 100.10


def test_export_requires_read_authorization(
    client: TestClient, world: World, auth: Callable[..., dict[str, str]]
) -> None:
    claimant = auth(world.claimant)
    outsider = auth(world.outsider)
    claim_id = _draft_with_line(client, claimant, world)

    for fmt in ("pdf", "xlsx"):
        response = client.get(f"/claims/{claim_id}/export/{fmt}", headers=outsider)
        assert response.status_code == 403, fmt
        assert response.json()["error"]["code"] == "authorization_denied"


def test_export_missing_claim_returns_404(
    client: TestClient, world: World, auth: Callable[..., dict[str, str]]
) -> None:
    response = client.get("/claims/99999999/export/pdf", headers=auth(world.claimant))
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "not_found"


def test_export_manager_can_read_team_claim(
    client: TestClient, world: World, auth: Callable[..., dict[str, str]]
) -> None:
    claim_id = _draft_with_line(client, auth(world.claimant), world)
    response = client.get(f"/claims/{claim_id}/export/pdf", headers=auth(world.manager))
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
