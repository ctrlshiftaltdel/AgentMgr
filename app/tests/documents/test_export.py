"""W1-10 报销单导出测试：PDF 与明细 XLSX。

覆盖任务第四步第 1–6 项：

1. PDF 字节以 ``b'%PDF-'`` 开头且长度 > 1024；
2. ``claim_pdf_lines`` 含单号/员工名/币种，且合计金额精确为 ``'123.45'``；
3. 精度用例：明细 ``0.10`` 与 ``0.20`` 的合计格式化后为 ``'0.30'``；
4. XLSX 可用 ``openpyxl.load_workbook`` 重新加载；表头文本正确；明细行金额单元格
   为数值且 ``== 123.45``；合计行正确；
5. 中文可渲染：含中文公司名与描述的 PDF 生成不报错且长度 > 1024；
6. 纯函数断言：``export`` 模块不导入数据库/ORM 依赖，测试全程不需要数据库连接。

测试一律使用内存 ``BytesIO`` 与 ``tmp_path``，**绝不向仓库目录写入任何文件**。
"""

from __future__ import annotations

import io
import pathlib
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook

from app.documents import export

COMPANY = "北京示例科技有限公司"
EMPLOYEE = "张三"


def _claim(**overrides: object) -> SimpleNamespace:
    """构造一个报销单鸭子类型对象（不依赖 ORM）。"""
    base: dict[str, object] = {
        "claim_no": "BX2026100001",
        "currency": "CNY",
        "status": "approved",
        "total_amount": Decimal("123.45"),
    }
    base.update(overrides)
    return SimpleNamespace(**base)


def _line(**overrides: object) -> SimpleNamespace:
    """构造一条报销明细鸭子类型对象（不依赖 ORM）。"""
    base: dict[str, object] = {
        "expense_date": date(2026, 10, 1),
        "description": "差旅费",
        "amount": Decimal("123.45"),
        "tax_amount": Decimal("0.00"),
    }
    base.update(overrides)
    return SimpleNamespace(**base)


def _total_amount(text_lines: list[str]) -> str:
    """从合计行中取出金额文本（``'合计: 123.45'`` -> ``'123.45'``）。"""
    total_line = text_lines[-1]
    assert total_line.startswith("合计")
    return total_line.split(": ", 1)[1]


# --- 1. PDF 魔数与长度 ---------------------------------------------------------


def test_render_claim_pdf_starts_with_magic_and_is_non_trivial() -> None:
    """PDF 字节以 b'%PDF-' 开头且长度 > 1024。"""
    data = export.render_claim_pdf(
        _claim(), [_line()], employee_name=EMPLOYEE, company_name=COMPANY
    )
    assert data.startswith(b"%PDF-")
    assert len(data) > 1024


# --- 2. 文本行含关键字段 + 合计精确 -------------------------------------------


def test_claim_pdf_lines_contains_header_fields_and_exact_total() -> None:
    """文本行含单号/员工名/币种；合计金额精确为 '123.45'。"""
    text_lines = export.claim_pdf_lines(
        _claim(), [_line()], employee_name=EMPLOYEE, company_name=COMPANY
    )
    joined = "\n".join(text_lines)

    assert "BX2026100001" in joined
    assert EMPLOYEE in joined
    assert "CNY" in joined
    assert COMPANY in joined

    # 合计行金额精确等于 '123.45'，绝不是浮点产物。
    assert _total_amount(text_lines) == "123.45"
    assert "123.45000000000001" not in joined


# --- 3. 精度用例 ---------------------------------------------------------------


def test_decimal_tenths_sum_formats_exactly_as_thirty_cents() -> None:
    """明细 0.10 与 0.20 的合计格式化后为 '0.30'（不是 0.30000000000000004）。"""
    total = Decimal("0.10") + Decimal("0.20")  # Decimal 精确求和 == Decimal("0.30")
    text_lines = export.claim_pdf_lines(
        _claim(total_amount=total),
        [_line(amount=Decimal("0.10")), _line(amount=Decimal("0.20"))],
        employee_name="李四",
        company_name=COMPANY,
    )
    assert _total_amount(text_lines) == "0.30"
    assert "0.30000000000000004" not in "\n".join(text_lines)


def test_amount_is_formatted_with_decimal_not_float() -> None:
    """证明格式化未经过浮点：Decimal ROUND_HALF_UP 得 '0.13'，而浮点会得 '0.12'。"""
    text_lines = export.claim_pdf_lines(
        _claim(total_amount=Decimal("0.125")),
        [_line(amount=Decimal("0.125"))],
        employee_name="王五",
        company_name=COMPANY,
    )
    assert _total_amount(text_lines) == "0.13"


# --- 4. XLSX 往返 --------------------------------------------------------------


def test_render_claim_xlsx_roundtrips_with_numeric_amounts() -> None:
    """XLSX 可重新加载；表头正确；明细金额为数值且 == 123.45；合计行正确。"""
    data = export.render_claim_xlsx(_claim(), [_line()])

    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.worksheets[0]

    # 首行表头。
    headers = [cell.value for cell in sheet[1]]
    assert headers[:4] == ["日期", "描述", "金额", "税额"]

    # 明细行：金额单元格为数值且 == 123.45。
    detail = sheet[2]
    assert detail[0].value == "2026-10-01"
    assert detail[1].value == "差旅费"
    assert isinstance(detail[2].value, (int, float))
    assert detail[2].value == 123.45

    # 合计行：最后一行的金额 == 123.45。
    total_row = sheet[sheet.max_row]
    assert total_row[0].value == "合计"
    assert isinstance(total_row[2].value, (int, float))
    assert total_row[2].value == 123.45


def test_render_claim_xlsx_total_tax_is_decimal_sum() -> None:
    """XLSX 合计行的税额为明细税额的 Decimal 精确求和。"""
    data = export.render_claim_xlsx(
        _claim(total_amount=Decimal("0.30")),
        [
            _line(amount=Decimal("0.10"), tax_amount=Decimal("0.10")),
            _line(amount=Decimal("0.20"), tax_amount=Decimal("0.20")),
        ],
    )
    sheet = load_workbook(io.BytesIO(data)).active
    total_row = sheet[sheet.max_row]
    assert total_row[3].value == 0.30


# --- 5. 中文可渲染 -------------------------------------------------------------


def test_render_claim_pdf_renders_chinese_text() -> None:
    """含中文公司名与描述的 PDF 生成不报错且长度 > 1024。"""
    data = export.render_claim_pdf(
        _claim(),
        [_line(description="北京出差交通费")],
        employee_name=EMPLOYEE,
        company_name=COMPANY,
    )
    assert data.startswith(b"%PDF-")
    assert len(data) > 1024


# --- 6. 纯函数（不导入数据库/ORM） --------------------------------------------


def test_export_module_is_pure_and_has_no_db_imports() -> None:
    """export 源码不含数据库/ORM 引用；模块未拉入会话对象。"""
    source = pathlib.Path(export.__file__).read_text(encoding="utf-8")
    assert "app.db" not in source
    assert "sqlalchemy" not in source
    assert not hasattr(export, "session")
    # 全模块只有写入 XLSX 单元格边界处的一次浮点转换。
    assert source.count("float(") == 1
